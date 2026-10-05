import csv
import openpyxl
import pytest
import os

class DataReader:
    """A data utility class providing static methods to parse external CSV and Excel (.xlsx) files 
    into formatted datasets compatible with 'pytest.param' for data-driven testing vectors.
    """

    @staticmethod
    def excel_to_pytest_param(filepath, sheet_name="Sheet1", id_column=0):
        """Loads an Excel workbook, converts all row cell values into explicit strings, 
        and packages them into parameterized test data objects.

        Args:
            filepath (str): The absolute system path to the target .xlsx file.
            sheet_name (str): The specific worksheet name to read data from.
            id_column (int): The 0-based column index to use as the test case identifier.

        Returns:
            list: A list containing parameterized 'pytest.param' row objects with designated test IDs.
        """
        wb = openpyxl.load_workbook(filepath)
        sheet = wb[sheet_name]
        data = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            values = [str(v) if v is not None else "" for v in row]
            data.append(pytest.param(*values, id=str(values[id_column])))
        return data

    def csv_reader_for_inventory_test(filepath, row_id=0):
        """Parses a target CSV catalog file using a dictionary reader into uncast parameterized test rows.

        Args:
            filepath (str): The absolute system path to the target .csv file.
            row_id (int): The 0-based row value index to apply as the structural test identity.

        Returns:
            list: A list containing structured 'pytest.param' collection objects.
        """
        data = []
        with open(filepath, newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                values = list(row.values())
                data.append(pytest.param(*values, id=str(values[row_id])))
            return data

    def cart_adding_item_check_excel(filepath, sheet_name="Sheet1", row_id=0):
        """Loads an Excel file and extracts its unmodified row entities into parameters without strict string casting.

        Args:
            filepath (str): The absolute system path to the target .xlsx workbook file.
            sheet_name (str): The target sheet title string.
            row_id (int): The target item index used to identify test data blocks.

        Returns:
            list: A list of raw parsed 'pytest.param' matrices.
        """
        wb = openpyxl.load_workbook(filepath)
        sheet = wb[sheet_name]
        data = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            values = list(row)
            data.append(pytest.param(*values, id=str(values[row_id])))
        return data

    @staticmethod
    def csv_read_to_pytest(filepath, row_id=0):
        """A standardized internal utility method converting text-based CSV data lists into pytest parameter parameters.

        Args:
            filepath (str): The absolute system file path.
            row_id (int): The position matrix tracker defining the unique test block metadata ID.

        Returns:
            list: Parsed pytest data collections.
        """
        data = []
        with open(filepath, newline='', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                values = list(row.values())
                data.append(pytest.param(*values, id=str(values[row_id])))
            return data

    @staticmethod
    def excel_read_to_pytest(filepath, sheet_name="Sheet1", row_id=0):
        """A standardized internal utility method mapping row data matrices to a specific sheet context for test execution.

        Args:
            filepath (str): The target workbook storage path.
            sheet_name (str): The targeted worksheet tab label.
            row_id (int): The column position index applied for identifying tests.

        Returns:
            list: Aggregated parameter configurations.
        """
        wb = openpyxl.load_workbook(filepath)
        sheet = wb[sheet_name]
        data = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            values = list(row)
            data.append(pytest.param(*values, id=str(values[row_id])))
        return data

    @staticmethod
    def read_file(filepath):
        """Inspects an input file extension dynamically and automatically routes processing to either 
        the internal CSV handler or Excel data-mapping handler.

        Args:
            filepath (str): The target file destination tracking path.

        Returns:
            list: A ready-to-use array collection of test parameter conditions.

        Raises:
            ValueError: Thrown if an unmapped file extension string (anything outside .csv or .xlsx) is supplied.
        """
        filename, file_extension = os.path.splitext(filepath)
        ext = file_extension.lower()
        if file_extension == ".csv":
            return DataReader.csv_read_to_pytest(filepath)
            
        if file_extension == ".xlsx":
            return DataReader.excel_read_to_pytest(filepath)
        raise ValueError(f"File fotmat not Supported {file_extension}")
